#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Conversor Simplificado de Markdown para Word (.docx)
"""

from docx import Document
from docx.shared import Pt, Inches
import re

def convert_md_to_docx(md_file, docx_file):
    """Converte arquivo Markdown para Word de forma simples e robusta"""
    
    # Criar documento
    doc = Document()
    
    # Configurar margens
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Ler arquivo markdown
    try:
        with open(md_file, 'r', encoding='utf-8') as f:
            content = f.read()
    except Exception as e:
        print(f"❌ Erro ao ler arquivo: {e}")
        return False
    
    lines = content.split('\n')
    
    in_code_block = False
    in_table = False
    table_data = []
    
    for line in lines:
        line_stripped = line.strip()
        
        # Código em bloco
        if line_stripped.startswith('```'):
            in_code_block = not in_code_block
            if not in_code_block:
                doc.add_paragraph()  # Espaço após código
            continue
        
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style = 'No Spacing'
            if p.runs:
                p.runs[0].font.name = 'Courier New'
                p.runs[0].font.size = Pt(9)
            continue
        
        # Cabeçalhos
        if line_stripped.startswith('# '):
            doc.add_heading(line_stripped[2:], level=1)
        elif line_stripped.startswith('## '):
            doc.add_heading(line_stripped[3:], level=2)
        elif line_stripped.startswith('### '):
            doc.add_heading(line_stripped[4:], level=3)
        elif line_stripped.startswith('#### '):
            doc.add_heading(line_stripped[5:], level=4)
        
        # Tabelas
        elif line_stripped.startswith('|'):
            if not in_table:
                in_table = True
                table_data = []
            
            # Ignorar linha separadora
            if '---' in line_stripped:
                continue
            
            # Extrair células
            cells = [cell.strip() for cell in line_stripped.split('|')[1:-1]]
            if cells:
                table_data.append(cells)
        
        # Fim da tabela
        elif in_table and not line_stripped.startswith('|'):
            if len(table_data) > 1:
                # Criar tabela no Word
                num_cols = len(table_data[0])
                num_rows = len(table_data)
                
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.style = 'Light Grid Accent 1'
                
                for row_idx, row_data in enumerate(table_data):
                    for col_idx, cell_text in enumerate(row_data):
                        if col_idx < num_cols:
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = cell_text
                            # Negrito no cabeçalho
                            if row_idx == 0 and cell.paragraphs:
                                for run in cell.paragraphs[0].runs:
                                    run.bold = True
                
                doc.add_paragraph()  # Espaço após tabela
            
            in_table = False
            table_data = []
            
            # Processar linha atual
            if line_stripped:
                add_formatted_paragraph(doc, line_stripped)
        
        # Linhas normais
        elif not in_table:
            if line_stripped:
                add_formatted_paragraph(doc, line_stripped)
            else:
                doc.add_paragraph()  # Linha vazia
    
    # Salvar documento
    try:
        doc.save(docx_file)
        print(f"✅ Documento Word criado com sucesso!")
        print(f"📄 Local: {docx_file}")
        return True
    except Exception as e:
        print(f"❌ Erro ao salvar documento: {e}")
        return False

def add_formatted_paragraph(doc, text):
    """Adiciona parágrafo com formatação básica"""
    
    # Lista com marcadores
    if text.startswith('- ') or text.startswith('* '):
        content = text[2:]
        p = doc.add_paragraph(content, style='List Bullet')
    
    # Lista numerada
    elif re.match(r'^\d+\.\s', text):
        content = re.sub(r'^\d+\.\s', '', text)
        p = doc.add_paragraph(content, style='List Number')
    
    # Linha horizontal
    elif text.startswith('---'):
        doc.add_paragraph('_' * 80)
    
    # Parágrafo normal
    else:
        p = doc.add_paragraph()
        
        # Processar negrito **texto**
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part:
                p.add_run(part)

if __name__ == "__main__":
    md_file = r"C:\Users\miche\.gemini\antigravity\brain\9c43aece-c21b-4753-a0d7-72bdb569c623\examepad_audit_report.md"
    docx_file = r"C:\Users\miche\Downloads\ExamePad_Auditoria_Completa.docx"
    
    print("🔄 Convertendo Markdown para Word...")
    print(f"📖 Origem: {md_file}")
    print(f"💾 Destino: {docx_file}")
    print()
    
    success = convert_md_to_docx(md_file, docx_file)
    
    if success:
        print()
        print("🎉 Conversão concluída!")
        print("📂 Você pode abrir o arquivo em: C:\\Users\\miche\\Downloads\\")
    else:
        print()
        print("⚠️ Houve um problema na conversão.")
